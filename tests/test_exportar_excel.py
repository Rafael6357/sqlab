"""Spec exportar-excel — guardar el resultado en .xlsx real."""
from __future__ import annotations


def _exportar_xlsx(app, monkeypatch, destino):
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        lambda *a, **k: (str(destino), "Excel (*.xlsx)"),
    )
    app.exportar_resultado_excel()
    return destino


def test_xl01_boton_existe(app):
    """XL-01: botón EXPORTAR EXCEL junto al CSV."""
    assert hasattr(app, "btn_exportar_excel")
    assert app.btn_exportar_excel.text() == "EXPORTAR EXCEL"


def test_xl02_celdas_y_null(app, monkeypatch, tmp_path):
    """XL-02: cabecera + filas por celda; None → NULL; anchos razonables."""
    import openpyxl

    app._mostrar_resultado(["id", "nombre"], [[1, None], [2, "Ana Pérez"]])
    dest = _exportar_xlsx(app, monkeypatch, tmp_path / "res.xlsx")
    wb = openpyxl.load_workbook(str(dest), data_only=True)
    ws = wb.active
    assert [ws.cell(1, 1).value, ws.cell(1, 2).value] == ["id", "nombre"]
    assert [ws.cell(2, 1).value, ws.cell(2, 2).value] == [1, "NULL"]
    assert [ws.cell(3, 1).value, ws.cell(3, 2).value] == [2, "Ana Pérez"]
    assert ws.column_dimensions["B"].width >= len("Ana Pérez")
    wb.close()


def test_xl04_cancelar_y_vacio(app, monkeypatch):
    """XL-04: cancelar → sin fichero; sin resultados → toast, sin diálogo."""
    from PySide6.QtWidgets import QFileDialog
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: ("", ""))
    app._mostrar_resultado(["a"], [["1"]])
    antes = app.toast_msg
    app.exportar_resultado_excel()
    assert app.toast_msg == antes
    llamadas = []
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        lambda *a, **k: (llamadas.append(1), ("x",))[1],
    )
    app._limpiar_resultado()
    app._ultimo_resultado = ([], [])
    app.exportar_resultado_excel()
    assert llamadas == []
    assert "NADA QUE EXPORTAR" in app.toast_msg
