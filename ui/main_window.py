"""Ventana principal SQLab — terminal de práctica SQL.

100% offline con PySide6 y SQLite en memoria.
Sin CDN: solo fuentes monoespaciadas del sistema + QSS.
"""
from __future__ import annotations

import csv
import json
import os
import sys
import time
from datetime import datetime

from PySide6.QtCore import QEvent, QSettings, Qt, QTimer
from PySide6.QtGui import QAction, QCloseEvent, QIcon, QTextCursor
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QCompleter,
    QDialog,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QListWidgetItem,
    QMainWindow,
    QPlainTextEdit,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)

from core.session_loader import (
    Ejercicio,
    combinar_resultados,
    load_file,
    load_tablas_folder,
)
from core.sqlite_engine import SQLEngine, Table
from ui.crono import CronoMixin  # split Fase 2
from ui.dialogs import CLAUDE_PROMPT, _show_custom_dialog  # re-export compat (tests)
from ui.formato_sql import SQL_KEYWORDS, _formatear_sql  # re-export compat (tests)
from ui.paneles import PanelesMixin, _ruta_logo  # split Fase 2
from ui.sql_highlighter import SQLHighlighter
from ui.tablas import (  # split 3/3 + AE (auto-espaciado)
    _configurar_grilla_ancha,
    _item_grilla,
    ajustar_y_repartir,
    guardar_base_anchos,
    repartir_anchos_proporcional,
)

# Funciones SQL que el autocompletado cierra con () (spec autocompletar-parentesis)
_FUNCIONES_AUTOCIERRE = frozenset({"COUNT", "SUM", "AVG", "MIN", "MAX", "ROUND", "LENGTH", "COALESCE"})


class MainWindow(CronoMixin, PanelesMixin, QMainWindow):
    # Topes de rendimiento (spec rendimiento-tablas-grandes)
    VISOR_MAX_FILAS = 2000
    RESULTADO_MAX_FILAS = 5000
    MUESTRA_MEDICION = 100
    AVISO_MB = 50

    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("SQLab — Terminal de práctica SQL")
        self.setWindowIcon(QIcon(_ruta_logo()))
        self.resize(1360, 840)
        self.engine = SQLEngine()
        self.ejercicio = Ejercicio()
        self.historial: list[str] = []
        self._ultimo_resultado: tuple[list[str], list[list]] = ([], [])
        self.settings = QSettings("SQLPractica", "SQLPractica")

        self._completer: QCompleter | None = None
        self.crono_activo = False
        self.crono_alerta = False
        self._crono_acumulado = 0.0
        self._crono_base = 0.0
        self._build_ui()
        self._connect_signals()
        self._apply_settings()
        self._setup_shortcuts()
        self._start_timers()
        self._load_initial_preset()

    # ------------------------------------------------------------------ UI

    def _build_ui(self) -> None:
        central = QWidget()
        root = QVBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        root.addWidget(self._build_hud())
        root.addWidget(self._build_banner())
        root.addWidget(self._build_hint_drawer())

        main_splitter = QSplitter(Qt.Orientation.Horizontal)
        main_splitter.setHandleWidth(1)
        main_splitter.addWidget(self._build_matrix())
        main_splitter.addWidget(self._build_workspace())
        main_splitter.setStretchFactor(0, 0)
        main_splitter.setStretchFactor(1, 1)
        main_splitter.setSizes([288, 1072])
        root.addWidget(main_splitter, stretch=1)
        self.setCentralWidget(central)

    def _build_workspace(self) -> QWidget:
        wrap = QWidget()
        lay = QVBoxLayout(wrap)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)

        top = QFrame()
        top.setObjectName("TopPanel")
        top_lay = QVBoxLayout(top)
        top_lay.setContentsMargins(0, 0, 0, 0)
        top_lay.setSpacing(0)
        tabhead = QHBoxLayout()
        tabhead.setContentsMargins(10, 0, 10, 0)
        tabhead.setSpacing(4)
        self.top_tabs = QTabWidget()
        self.top_tabs.addTab(self._build_mission_tab(), "[EJERCICIO]")
        self.top_tabs.setTabToolTip(0, "Enunciado del ejercicio y columnas que debe devolver tu consulta")
        self.top_tabs.addTab(self._build_dump_tab(), "CONTENIDO DE LA TABLA")
        self.top_tabs.setTabToolTip(1, "Datos de la tabla activa (solo lectura)")
        tabhead.addWidget(self.top_tabs, stretch=1)
        top_lay.addLayout(tabhead)
        lay.addWidget(top, stretch=4)

        bottom = QFrame()
        bottom.setObjectName("BottomPanel")
        bl = QVBoxLayout(bottom)
        bl.setContentsMargins(0, 0, 0, 0)
        bl.setSpacing(0)
        bl.addWidget(self._build_deck())
        self.editor_pane = self._build_editor_pane()
        self.output_pane = self._build_output_pane()
        work = QSplitter(Qt.Orientation.Horizontal)
        work.setHandleWidth(6)
        work.setChildrenCollapsible(False)
        work.addWidget(self.editor_pane)
        work.addWidget(self.output_pane)
        work.setCollapsible(0, False)
        work.setCollapsible(1, False)
        work.setStretchFactor(0, 1)
        work.setStretchFactor(1, 1)
        work.setSizes([500, 500])
        self.work_splitter = work
        # MR-02: doble-clic en el handle resetea a ~50/50
        try:
            h = work.handle(1)
            if h is not None:
                h.installEventFilter(self)
        except Exception:
            pass
        bl.addWidget(work, stretch=1)
        lay.addWidget(bottom, stretch=6)
        return wrap

    def _build_mission_tab(self) -> QWidget:
        page = QWidget()
        lay = QHBoxLayout(page)
        lay.setContentsMargins(10, 8, 10, 8)
        lay.setSpacing(8)

        spec = QFrame()
        spec.setObjectName("SpecCard")
        sl = QVBoxLayout(spec)
        sl.setContentsMargins(10, 8, 10, 8)
        sl.setSpacing(4)
        head = QHBoxLayout()
        head.setContentsMargins(0, 0, 0, 0)
        spec_t = QLabel("■ ESPECIFICACIÓN DE CONSULTA")
        spec_t.setObjectName("PanelTitle")
        head.addWidget(spec_t)
        head.addStretch()
        self.btn_copiar_spec = QPushButton("COPIAR")
        self.btn_copiar_spec.setObjectName("GhostBtn")
        self.btn_copiar_spec.setToolTip("Copiar la especificación (título + enunciado + columnas objetivo)")
        self.btn_copiar_spec.clicked.connect(self.copiar_especificacion)
        head.addWidget(self.btn_copiar_spec)
        oid = QLabel("OBJETIVO N.º 001")
        oid.setObjectName("MutedLabel")
        head.addWidget(oid)
        sl.addLayout(head)
        self.enunciado_texto = QLabel("Carga un ejercicio .json para empezar.")
        self.enunciado_texto.setObjectName("StatementText")
        self.enunciado_texto.setWordWrap(True)
        sl.addWidget(self.enunciado_texto)
        sl.addStretch()
        tgt = QLabel("COLUMNAS OBJETIVO:")
        tgt.setObjectName("MutedLabel")
        sl.addWidget(tgt)
        chips = QHBoxLayout()
        chips.setContentsMargins(0, 0, 0, 0)
        self.target_chip1 = QLabel("—")
        self.target_chip1.setObjectName("TargetColGreen")
        chips.addWidget(self.target_chip1)
        plus = QLabel("+")
        plus.setObjectName("MutedLabel")
        chips.addWidget(plus)
        self.target_chip2 = QLabel("—")
        self.target_chip2.setObjectName("TargetColCyan")
        chips.addWidget(self.target_chip2)
        chips.addStretch()
        self.sort_label = QLabel("")
        self.sort_label.setObjectName("MutedLabel")
        chips.addWidget(self.sort_label)
        sl.addLayout(chips)
        self.expected_cols = QLabel("")  # compat
        self.expected_cols.setVisible(False)
        sl.addWidget(self.expected_cols)
        lay.addWidget(spec, stretch=1)
        return page

    def _build_dump_tab(self) -> QWidget:
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setContentsMargins(0, 0, 0, 0)
        self.visor_tabla = QTableWidget()
        self.visor_tabla.setObjectName("ResultTable")
        self.visor_tabla.setToolTip("Datos de la tabla activa (solo lectura)")
        self.visor_tabla.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.visor_tabla.setAlternatingRowColors(True)
        _configurar_grilla_ancha(self.visor_tabla)
        lay.addWidget(self.visor_tabla)
        return page

    def _build_deck(self) -> QFrame:
        bar = QFrame()
        bar.setObjectName("ExerciseBanner")
        bar.setFixedHeight(36)
        lay = QHBoxLayout(bar)
        lay.setContentsMargins(10, 0, 10, 0)
        lay.setSpacing(8)
        title = QLabel("> CONSOLA SQL")
        title.setObjectName("PanelTitle")
        lay.addWidget(title)
        lay.addWidget(self._sep())
        hot = QLabel("ATAJO: CTRL + ENTER PARA EJECUTAR")
        hot.setObjectName("MutedLabel")
        lay.addWidget(hot)
        lay.addStretch()
        self.autocomplete_check = QCheckBox("AUTOCOMPLETAR")
        self.autocomplete_check.setToolTip("Autocompletado: sugiere tablas, columnas y palabras clave. Apagado por defecto.")
        lay.addWidget(self.autocomplete_check)
        self.btn_format = QPushButton("FORMATO SQL")
        self.btn_format.setObjectName("FormatBtn")
        self.btn_format.setToolTip("Aplica formato SQL estándar: mayúsculas, saltos por cláusula e indentación")
        self.btn_clear_editor = QPushButton("✕")
        self.btn_clear_editor.setObjectName("GhostBtn")
        self.btn_clear_editor.setFixedWidth(30)
        self.btn_clear_editor.setToolTip("Limpiar editor")
        self.btn_copiar = QPushButton("COPIAR PARA IA")
        self.btn_copiar.setObjectName("CyberBtn")
        self.btn_copiar.setToolTip("Copiar la consulta con plantilla lista para pegar en IA")
        self.btn_ejecutar = QPushButton("EJECUTAR_SQL")
        self.btn_ejecutar.setObjectName("ExecuteBtn")
        self.btn_ejecutar.setToolTip("Ejecutar (F5 o Ctrl+Enter)")
        for b in (self.btn_format, self.btn_clear_editor, self.btn_copiar, self.btn_ejecutar):
            lay.addWidget(b)
        return bar

    def _build_editor_pane(self) -> QFrame:
        pane = QFrame()
        pane.setObjectName("EditorPane")
        pane.setMinimumWidth(220)
        lay = QVBoxLayout(pane)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)
        self.editor = QPlainTextEdit()
        self.editor.setObjectName("SQLEditor")
        self.editor.setToolTip("Escribe tu consulta SQL aquí (Ctrl+Enter para ejecutar)")
        self.editor.setPlaceholderText(
            "-- ESCRIBE TU CONSULTA SQL AQUÍ...\nSELECT * FROM clientes;"
        )
        self.editor.setTabChangesFocus(False)
        self.highlighter = SQLHighlighter(self.editor.document())
        # AC-01/AC-02: el filtro acepta la sugerencia con Enter/Tab (ver eventFilter)
        self.editor.installEventFilter(self)
        lay.addWidget(self.editor, stretch=1)
        status = QFrame()
        status.setObjectName("ExerciseBanner")
        status.setFixedHeight(24)
        sl = QHBoxLayout(status)
        sl.setContentsMargins(10, 0, 10, 0)
        sl.setSpacing(6)
        self.status_dot = QLabel("●")
        self.status_dot.setObjectName("PanelTitle")
        sl.addWidget(self.status_dot)
        self.cursor_label = QLabel("LÍN 1, COL 1")
        self.cursor_label.setObjectName("StatusLabel")
        sl.addWidget(self.cursor_label)
        sl.addStretch()
        d = QLabel("DIALECTO: SQLITE3")
        d.setObjectName("StatusLabel")
        sl.addWidget(d)
        u = QLabel("UTF-8 // CRLF")
        u.setObjectName("StatusLabel")
        sl.addWidget(u)
        lay.addWidget(status)
        return pane

    def _build_output_pane(self) -> QFrame:
        pane = QFrame()
        pane.setObjectName("ResultPane")
        pane.setMinimumWidth(240)
        lay = QVBoxLayout(pane)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(0)
        head = QFrame()
        head.setObjectName("ExerciseBanner")
        head.setFixedHeight(28)
        hl = QHBoxLayout(head)
        hl.setContentsMargins(10, 0, 10, 0)
        hl.setSpacing(8)
        t = QLabel(">> RESULTADO DE LA CONSULTA")
        t.setObjectName("PanelTitle")
        hl.addWidget(t)
        self.row_badge = QLabel("0 FILAS")
        self.row_badge.setObjectName("RowBadge")
        self.row_badge.setVisible(False)
        hl.addWidget(self.row_badge)
        self.btn_exportar = QPushButton("EXPORTAR CSV")
        self.btn_exportar.setObjectName("GhostBtn")
        self.btn_exportar.setToolTip("Guardar el resultado completo en un .csv (UTF-8, abre en Excel)")
        self.btn_exportar.clicked.connect(self.exportar_resultado_csv)
        hl.addWidget(self.btn_exportar)
        self.btn_exportar_excel = QPushButton("EXPORTAR EXCEL")
        self.btn_exportar_excel.setObjectName("GhostBtn")
        self.btn_exportar_excel.setToolTip("Guardar el resultado completo en un .xlsx (celdas separadas, abre en Excel)")
        self.btn_exportar_excel.clicked.connect(self.exportar_resultado_excel)
        hl.addWidget(self.btn_exportar_excel)
        self.btn_graficar = QPushButton("GRAFICAR")
        self.btn_graficar.setObjectName("GhostBtn")
        self.btn_graficar.setToolTip("Graficar el resultado (2 columnas: etiqueta, valor)")
        self.btn_graficar.clicked.connect(self.graficar_resultado)
        hl.addWidget(self.btn_graficar)
        hl.addStretch()
        self.exec_time = QLabel("EN ESPERA")
        self.exec_time.setObjectName("StatusLabel")
        hl.addWidget(self.exec_time)
        lay.addWidget(head)

        self.empty_state = QLabel("EN ESPERA // LISTO PARA EJECUTAR")
        self.empty_state.setObjectName("MutedLabel")
        self.empty_state.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(self.empty_state, stretch=1)

        self.error_box = QFrame()
        self.error_box.setObjectName("ErrorBox")
        el = QVBoxLayout(self.error_box)
        el.setContentsMargins(10, 8, 10, 8)
        el.setSpacing(4)
        et = QHBoxLayout()
        et.setContentsMargins(0, 0, 0, 0)
        err_title = QLabel("EXCEPCIÓN_SINTAXIS_SQLITE")
        err_title.setObjectName("PanelTitle")
        et.addWidget(err_title)
        et.addStretch()
        code = QLabel("CÓD_ERROR: 0x22")
        code.setObjectName("MutedLabel")
        et.addWidget(code)
        el.addLayout(et)
        self.error_text = QLabel("")
        self.error_text.setWordWrap(True)
        el.addWidget(self.error_text)
        adv = QHBoxLayout()
        adv.setContentsMargins(0, 0, 0, 0)
        adv_t = QLabel(">> CONSEJO DE RECUPERACIÓN:")
        adv_t.setObjectName("HintAccent")
        adv.addWidget(adv_t)
        self.error_hint = QLabel("")
        self.error_hint.setObjectName("MutedLabel")
        self.error_hint.setWordWrap(True)
        adv.addWidget(self.error_hint, stretch=1)
        el.addLayout(adv)
        self.error_box.setVisible(False)
        lay.addWidget(self.error_box)

        self.resultado_tabla = QTableWidget()
        self.resultado_tabla.setObjectName("ResultTable")
        self.resultado_tabla.setToolTip("Resultados de la consulta (solo lectura)")
        _configurar_grilla_ancha(self.resultado_tabla)
        self.resultado_tabla.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.resultado_tabla.setAlternatingRowColors(True)
        self.resultado_tabla.setVisible(False)
        lay.addWidget(self.resultado_tabla, stretch=1)

        self.mensaje_label = QLabel("")
        self.mensaje_label.setObjectName("MutedLabel")
        self.mensaje_label.setVisible(False)
        lay.addWidget(self.mensaje_label)
        # compat extra
        self.tabVisualTableName = QLabel("")
        self.tabVisualTableName.setVisible(False)
        self.toast_msg = ""
        return pane

    # ------------------------------------------------------------- signals

    def _connect_signals(self) -> None:
        self.btn_cargar.clicked.connect(self.cargar_unificado)
        self.btn_save.clicked.connect(self.guardar_sesion)
        self.btn_exportar_db.clicked.connect(self.exportar_db)
        self.btn_ses.clicked.connect(self.cargar_sesion)
        self.btn_json.clicked.connect(self.mostrar_formato_json)
        self.btn_ejecutar.clicked.connect(self.ejecutar_consulta)
        self.btn_copiar.clicked.connect(self.copiar_consulta)
        self.btn_preset_tienda.clicked.connect(lambda: self.cargar_preset("ejemplo_tienda.json"))
        self.btn_preset_biblio.clicked.connect(lambda: self.cargar_preset("ejemplo_biblioteca.json"))
        self.btn_reset.clicked.connect(self.restablecer_datos)
        self.btn_select_all.clicked.connect(self.ver_select_all)
        self.btn_clear_hist.clicked.connect(self.limpiar_historial)
        self.btn_comparar.clicked.connect(self.comparar_desde_historial)
        self.btn_format.clicked.connect(self.formatear_consulta)
        self.btn_clear_editor.clicked.connect(self.limpiar_editor)
        self.tabla_list.currentItemChanged.connect(self._on_tabla_selected)
        self.columnas_list.itemClicked.connect(self._on_columna_clicked)
        self.pista_toggle.toggled.connect(self._on_pista_toggle)
        self.historial_toggle.toggled.connect(self._on_historial_toggle)
        self.autocomplete_check.toggled.connect(self._on_autocomplete_toggle)
        self.historial_list.itemClicked.connect(self._on_historial_clicked)
        self.editor.textChanged.connect(self._on_editor_text_changed)
        # AE-02/AE-03: reparto proporcional al estirar; reajuste manual = nueva base
        self._grillas_auto = (self.visor_tabla, self.resultado_tabla)
        for grilla in self._grillas_auto:
            grilla.installEventFilter(self)
            grilla.horizontalHeader().sectionResized.connect(
                lambda _l, _o, _n, g=grilla: guardar_base_anchos(g)
            )
        self.editor.cursorPositionChanged.connect(self._update_cursor_pos)

    def _setup_shortcuts(self) -> None:
        act_f5 = QAction(self)
        act_f5.setShortcut("F5")
        act_f5.triggered.connect(self.ejecutar_consulta)
        self.addAction(act_f5)
        act_ctrl_enter = QAction(self)
        act_ctrl_enter.setShortcut("Ctrl+Return")
        act_ctrl_enter.triggered.connect(self.ejecutar_consulta)
        self.addAction(act_ctrl_enter)

    def _start_timers(self) -> None:
        self.blink = QTimer(self)
        self.blink.timeout.connect(self._blink_dot)
        self.blink.start(1000)
        self._dot_on = True
        self._crono_timer = QTimer(self)
        self._crono_timer.setInterval(1000)
        self._crono_timer.timeout.connect(self._crono_tick)

    def _blink_dot(self) -> None:
        self._dot_on = not self._dot_on
        self.status_dot.setText("●" if self._dot_on else "○")

    # --------------------------------------------------- splitter reset (MR-02)
    def _reset_work_splitter(self) -> None:
        if not hasattr(self, "work_splitter"):
            return
        w = self.work_splitter.width() or 1000
        self.work_splitter.setSizes([w // 2, w - w // 2])

    def eventFilter(self, obj, event):  # type: ignore[override]
        try:
            if hasattr(self, "work_splitter") and obj is self.work_splitter.handle(1):
                if event.type() == QEvent.Type.MouseButtonDblClick:
                    self._reset_work_splitter()
                    return True
            # AC-01/AC-02/AC-03: Enter/Tab acepta la sugerencia, Escape la cierra.
            # Ctrl+Enter se deja pasar (lo ejecuta el atajo Ctrl+Return).
            if obj is getattr(self, "editor", None) and event.type() == QEvent.Type.KeyPress:
                tecla = event.key()
                if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
                    return super().eventFilter(obj, event)
                if tecla in (Qt.Key.Key_Enter, Qt.Key.Key_Return, Qt.Key.Key_Tab):
                    if self._aceptar_autocompletado():
                        return True
                elif tecla == Qt.Key.Key_ParenLeft:
                    # PR-03: ( manual se autocierra solo sin popup visible
                    completer = getattr(self, "_completer", None)
                    if completer is not None and completer.popup().isVisible():
                        return super().eventFilter(obj, event)
                    tc = self.editor.textCursor()
                    if tc.hasSelection():
                        sel = tc.selectedText()
                        tc.insertText(f"({sel})")
                        tc.movePosition(QTextCursor.MoveOperation.Left)
                    else:
                        tc.insertText("()")
                        tc.movePosition(QTextCursor.MoveOperation.Left)
                    self.editor.setTextCursor(tc)
                    return True
                elif tecla == Qt.Key.Key_Escape:
                    completer = getattr(self, "_completer", None)
                    if completer is not None and completer.popup().isVisible():
                        completer.popup().hide()
                        return True
            # AE-02: al redimensionar visor/resultado, repartir proporcional
            if event.type() == QEvent.Type.Resize and obj in getattr(self, "_grillas_auto", ()):
                repartir_anchos_proporcional(obj)
                return super().eventFilter(obj, event)
        except Exception:
            pass
        return super().eventFilter(obj, event)

    def _aceptar_autocompletado(self) -> bool:
        """Inserta la sugerencia resaltada (o la primera) si el popup está visible."""
        completer = getattr(self, "_completer", None)
        if completer is None or not completer.popup().isVisible():
            return False
        idx = completer.popup().currentIndex()
        if not idx.isValid():
            try:
                idx = completer.completionModel().index(0, 0)
            except Exception:
                return False
        if not idx.isValid():
            return False
        try:
            texto = completer.completionModel().data(idx, Qt.ItemDataRole.DisplayRole)
        except Exception:
            return False
        if not texto:
            return False
        completer.popup().hide()
        self._insert_completion(str(texto))
        return True

    # ------------------------------------------------------------- settings

    def _apply_settings(self) -> None:
        val = self.settings.value("autocompletado", "false")
        self.autocomplete_check.setChecked(str(val).lower() == "true")

    def _on_autocomplete_toggle(self, checked: bool) -> None:
        self.settings.setValue("autocompletado", "true" if checked else "false")
        if checked:
            self._completer = self._build_completer()
            if self._completer:
                self._completer.setWidget(self.editor)
                self._completer.activated.connect(self._insert_completion)
        else:
            if self._completer:
                try:
                    self._completer.activated.disconnect(self._insert_completion)
                except Exception:
                    pass
                self._completer = None

    def _build_completer(self) -> QCompleter | None:
        words = set(SQL_KEYWORDS)
        for t in self.engine.tables.values():
            words.add(t.name)
            words.update(c.name for c in t.columns)
        if not words:
            return None
        from PySide6.QtCore import QStringListModel
        model = QStringListModel(sorted(words), self)
        completer = QCompleter(model, self)
        completer.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        completer.setFilterMode(Qt.MatchFlag.MatchContains)
        completer.setWidget(self.editor)
        completer.setCompletionMode(QCompleter.CompletionMode.PopupCompletion)
        return completer

    def _text_under_cursor(self) -> str:
        tc = self.editor.textCursor()
        tc.select(QTextCursor.SelectionType.WordUnderCursor)
        return tc.selectedText()

    def _on_editor_text_changed(self) -> None:
        if not self.autocomplete_check.isChecked() or not self._completer:
            return
        prefix = self._text_under_cursor()
        if len(prefix) < 2:
            self._completer.popup().hide()
            return
        self._completer.setCompletionPrefix(prefix)
        if self._completer.completionCount() == 0:
            self._completer.popup().hide()
            return
        cr = self.editor.cursorRect()
        cr.setWidth(self._completer.popup().sizeHintForColumn(0) + 16)
        self._completer.complete(cr)
        # AC-01: resaltar la primera coincidencia para que Enter la acepte
        try:
            self._completer.popup().setCurrentIndex(self._completer.completionModel().index(0, 0))
        except Exception:
            pass

    def _insert_completion(self, text: str) -> None:
        tc = self.editor.textCursor()
        prefix = self._text_under_cursor()
        if prefix:
            tc.movePosition(QTextCursor.MoveOperation.Left, QTextCursor.MoveMode.KeepAnchor, len(prefix))
        tc.insertText(text)
        # PR-01: las funciones se completan con () y el cursor dentro
        if text.upper() in _FUNCIONES_AUTOCIERRE:
            tc.insertText("()")
            tc.movePosition(QTextCursor.MoveOperation.Left)
        self.editor.setTextCursor(tc)

    def _update_cursor_pos(self) -> None:
        tc = self.editor.textCursor()
        self.cursor_label.setText(f"LÍN {tc.blockNumber() + 1}, COL {tc.columnNumber() + 1}")

    # ------------------------------------------------------------ carga

    def _base_dir(self) -> str:
        return os.path.dirname(os.path.abspath(__file__))

    @staticmethod
    def _bundle_dir() -> str:
        """Raíz del bundle: _MEIPASS en exe frozen, raíz del repo en dev.

        PyInstaller onefile extrae `datas` (resources/, examples/) a
        sys._MEIPASS; en dev los recursos viven junto al código.
        """
        if getattr(sys, "frozen", False):
            meipass = getattr(sys, "_MEIPASS", None)
            if meipass:
                return meipass
        return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    def _load_initial_preset(self) -> None:
        # CV-01: al arrancar la consola queda vacía (sin defaultQuery inyectada)
        if not self.cargar_preset("ejemplo_tienda.json", silencioso=True, escribir_query=False):
            self._set_default_query()

    def _default_query_for(self) -> str | None:
        dq = (self.ejercicio.default_query or "").strip()
        if dq:
            return dq if dq.endswith(";") else dq + ";"
        names = self.engine.table_names()
        if names:
            return f"SELECT * FROM {names[0]};"
        return None

    def _set_default_query(self) -> None:
        q = self._default_query_for()
        if q:
            self.editor.setPlainText(q + "\n")

    def cargar_preset(self, filename: str, silencioso: bool = False, escribir_query: bool = True) -> bool:
        path = os.path.normpath(os.path.join(self._bundle_dir(), "examples", filename))
        if not os.path.exists(path):
            if not silencioso:
                _show_custom_dialog(self, "EJEMPLO NO ENCONTRADO", f"No se encontró:\n{path}")
            return False
        result = load_file(path)
        if not result.ok and not silencioso:
            _show_custom_dialog(self, "ERROR DE DECODIFICACIÓN", "\n".join(result.errors))
        self._aplicar_resultado(result, path, escribir_query=escribir_query)
        return result.ok

    def cargar_json(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "CARGAR EJERCICIO (.json)", "", "JSON (*.json)")
        if not path:
            return
        if not self._confirmar_archivo_grande([path]):
            return
        self._aplicar_resultado(load_file(path), path)

    def _elegir_modo_carga(self) -> str | None:
        """Mini-diálogo custom: EJERCICIO (.json), ARCHIVOS o CARPETA (spec cargar-unificado)."""
        dlg = QDialog(self)
        dlg.setWindowTitle("CARGAR")
        dlg.setMinimumWidth(460)
        dlg.setModal(True)
        lay = QVBoxLayout(dlg)
        lay.setContentsMargins(16, 16, 16, 16)
        lay.setSpacing(12)
        title_lbl = QLabel("CARGAR")
        title_lbl.setObjectName("PanelTitle")
        lay.addWidget(title_lbl)
        msg_lbl = QLabel("¿Qué quieres cargar? (.json, *.csv, *.xlsx, *.xls, *.db)")
        msg_lbl.setWordWrap(True)
        msg_lbl.setObjectName("StatementText")
        lay.addWidget(msg_lbl)
        row = QHBoxLayout()
        row.addStretch()
        eleccion: list[str | None] = [None]
        btn_ej = QPushButton("EJERCICIO")
        btn_ej.setObjectName("PrimaryBtn")
        btn_ej.setToolTip("Cargar un ejercicio desde un archivo .json (formato clásico o IA)")
        btn_ej.clicked.connect(lambda: (eleccion.__setitem__(0, "ejercicio"), dlg.accept()))
        btn_files = QPushButton("ARCHIVOS")
        btn_files.setObjectName("PrimaryBtn")
        btn_files.clicked.connect(lambda: (eleccion.__setitem__(0, "archivos"), dlg.accept()))
        btn_folder = QPushButton("CARPETA")
        btn_folder.setObjectName("GhostBtn")
        btn_folder.clicked.connect(lambda: (eleccion.__setitem__(0, "carpeta"), dlg.accept()))
        btn_cancel = QPushButton("CANCELAR")
        btn_cancel.setObjectName("GhostBtn")
        btn_cancel.clicked.connect(dlg.reject)
        row.addWidget(btn_ej)
        row.addWidget(btn_files)
        row.addWidget(btn_folder)
        row.addWidget(btn_cancel)
        lay.addLayout(row)
        dlg.exec()
        return eleccion[0]

    def cargar_unificado(self) -> None:
        """Entrada del botón CARGAR: despacha según el origen elegido (CU-02)."""
        modo = self._elegir_modo_carga()
        if modo == "ejercicio":
            self.cargar_json()
        elif modo == "archivos":
            self.cargar_tablas_archivos()
        elif modo == "carpeta":
            self.cargar_csv_carpeta()

    def cargar_tablas(self) -> None:
        """Compat: despacha a archivos o carpeta (ignora modo ejercicio)."""
        modo = self._elegir_modo_carga()
        if modo == "archivos":
            self.cargar_tablas_archivos()
        elif modo == "carpeta":
            self.cargar_csv_carpeta()

    def _confirmar_archivo_grande(self, paths: list[str]) -> bool:
        """Pre-aviso RG-05: si los ficheros superan AVISO_MB pide confirmación."""
        try:
            mb = sum(os.path.getsize(p) for p in paths) / (1024 * 1024)
        except OSError:
            return True
        if mb <= self.AVISO_MB:
            return True
        dlg = QDialog(self)
        dlg.setWindowTitle("ARCHIVO GRANDE")
        dlg.setMinimumWidth(420)
        dlg.setModal(True)
        lay = QVBoxLayout(dlg)
        lay.setContentsMargins(16, 16, 16, 16)
        lay.setSpacing(12)
        title_lbl = QLabel("ARCHIVO GRANDE")
        title_lbl.setObjectName("PanelTitle")
        lay.addWidget(title_lbl)
        msg_lbl = QLabel(
            f"Vas a cargar {mb:.0f} MB. Puede tardar varios minutos "
            f"y usar mucha memoria (tope de vista: {self.VISOR_MAX_FILAS} filas)."
        )
        msg_lbl.setWordWrap(True)
        msg_lbl.setObjectName("StatementText")
        lay.addWidget(msg_lbl)
        row = QHBoxLayout()
        row.addStretch()
        ok = QPushButton("CARGAR")
        ok.setObjectName("PrimaryBtn")
        ok.clicked.connect(dlg.accept)
        no = QPushButton("CANCELAR")
        no.setObjectName("GhostBtn")
        no.clicked.connect(dlg.reject)
        row.addWidget(ok)
        row.addWidget(no)
        lay.addLayout(row)
        return dlg.exec() == QDialog.DialogCode.Accepted

    def cargar_tablas_archivos(self) -> None:
        files, _ = QFileDialog.getOpenFileNames(
            self,
            "CARGAR TABLAS — archivos *.csv / *.xlsx / *.xls / *.db (UTF-8)",
            "",
            "Tablas (*.csv *.xlsx *.xls *.db *.sqlite *.sqlite3);;Todos los archivos (*.*)",
        )
        if not files:
            return
        if not self._confirmar_archivo_grande(sorted(files)):
            return
        resultados = [load_file(f) for f in sorted(files)]
        merged, reemplazadas = combinar_resultados(resultados)
        self._aplicar_resultado(merged, "; ".join(files))
        if merged.ok and reemplazadas:
            self._toast(
                f"TABLA(S) REEMPLAZADA(S): {', '.join(reemplazadas)} — GANÓ EL ÚLTIMO ARCHIVO"
            )

    def cargar_csv_carpeta(self) -> None:
        folder = QFileDialog.getExistingDirectory(
            self, "CARGAR TABLAS — carpeta con *.csv / *.xlsx / *.xls / *.db (UTF-8)"
        )
        if not folder:
            return
        archivos = [
            os.path.join(folder, f)
            for f in os.listdir(folder)
            if os.path.splitext(f)[1].lower() in (".csv", ".xlsx", ".xls", ".db", ".sqlite", ".sqlite3")
        ]
        if not self._confirmar_archivo_grande(archivos):
            return
        self._aplicar_resultado(load_tablas_folder(folder), folder)

    # Alias nuevo (compat UI / tests)
    def cargar_tablas_carpeta(self) -> None:
        return self.cargar_csv_carpeta()

    def _refresh_status(self) -> None:
        """Barra de estado real: conteos del engine (UN-01/UN-02)."""
        tablas = self.engine.table_names()
        filas = sum(len(t.rows) for t in self.engine.tables.values())
        self.status_db.setText(f"TABLAS: {len(tablas)} · FILAS: {filas} · DB: MEMORIA OK")

    def _aplicar_resultado(self, result, _origen: str, escribir_query: bool = True) -> None:
        if not result.ok:
            msg = "\n".join(result.errors) or "No se pudieron cargar las tablas."
            _show_custom_dialog(self, "ERROR DE DECODIFICACIÓN", msg)
            return
        self._crono_reset()  # CR-05: ejercicio nuevo → crono a cero y detenido
        omitidas = self.engine.load_tables(result.tables) or []
        self.ejercicio = result.ejercicio or Ejercicio()
        self._refresh_tabla_list()
        self._refresh_briefing()
        self._refresh_autocomplete()
        self._refresh_status()
        filas = sum(len(t.rows) for t in result.tables)
        self._toast(f"EJERCICIO CARGADO: {len(result.tables)} TABLA(S), {filas} FILA(S)")
        if escribir_query:
            self._set_default_query()
        self._limpiar_resultado()
        if omitidas:
            # CR-03: el engine descartó tablas (CREATE fallido) → avisar
            self._toast(f"TABLA(S) OMITIDA(S): {', '.join(omitidas)}")
        if result.errors:
            # CA-04: carga parcial (algunos ficheros fallaron) → avisar con nombres
            _show_custom_dialog(self, "ERROR DE DECODIFICACIÓN", "\n".join(result.errors))

    # ------------------------------------------------- matrix / briefing

    def _level_label(self) -> str:
        dif = (self.ejercicio.dificultad or "Principiante").strip()
        return dif.upper() if dif else "PRINCIPIANTE"

    def _expected_columns(self) -> list[str]:
        titulo = self.ejercicio.titulo or ""
        if "Top Clientes" in titulo or "Gasto" in titulo:
            return ["nombre", "total_gastado"]
        if "Libros" in titulo:
            return ["titulo", "lector", "dias_prestamo"]
        return []

    def _refresh_tabla_list(self) -> None:
        self.tabla_list.clear()
        for name in self.engine.table_names():
            t = self.engine.tables[name]
            item = QListWidgetItem(f"> {name}  [{len(t.rows)}F]")
            item.setData(Qt.ItemDataRole.UserRole, name)
            self.tabla_list.addItem(item)
        self.tables_count.setText(f"({self.tabla_list.count()})")
        if self.tabla_list.count():
            self.tabla_list.setCurrentRow(0)

    def _refresh_briefing(self) -> None:
        titulo = self.ejercicio.titulo or "SIN EJERCICIO"
        self.exercise_title.setText(titulo.upper())
        self.enunciado_texto.setText(self.ejercicio.enunciado or "Carga un ejercicio .json para empezar.")
        self.difficulty_badge.setText(self._level_label())
        cols = self._expected_columns()
        self.target_chip1.setText(cols[0] if len(cols) > 0 else "—")
        self.target_chip2.setText(cols[1] if len(cols) > 1 else "—")
        self.expected_cols.setText(", ".join(cols))
        if "Top Clientes" in titulo or "Gasto" in titulo:
            self.sort_label.setText("ORDEN: DESCENDENTE (ORDER BY total_gastado DESC)")
        else:
            self.sort_label.setText("")
        if self.ejercicio.pista:
            self.pista_label.setText(self.ejercicio.pista)
            self.pista_toggle.setVisible(True)
            self.pista_toggle.setChecked(False)
            self.pista_card.setVisible(False)
            self.pista_toggle.setText("VER_PISTA")
        else:
            self.pista_toggle.setVisible(False)
            self.pista_card.setVisible(False)

    def _refresh_autocomplete(self) -> None:
        if self.autocomplete_check.isChecked():
            self._on_autocomplete_toggle(True)

    def _on_pista_toggle(self, checked: bool) -> None:
        self.pista_card.setVisible(checked)
        self.pista_toggle.setText("OCULTAR_PISTA" if checked else "VER_PISTA")

    def _on_historial_toggle(self, checked: bool) -> None:
        self.historial_list.setVisible(checked)
        self.historial_toggle.setText("OCULTAR_HISTORIAL" if checked else "VER_HISTORIAL")

    def _on_tabla_selected(self, current: QListWidgetItem | None, _previous) -> None:
        if not current:
            self.columnas_list.clear()
            self.columnas_empty.setVisible(True)
            self.columnas_list.setVisible(False)
            return
        name = current.data(Qt.ItemDataRole.UserRole)
        table = self.engine.tables.get(name)
        if not table:
            return
        self.schema_label.setText(table.name)
        self.tabVisualTableName.setText(table.name)
        self.columnas_empty.setVisible(False)
        self.columnas_list.setVisible(True)
        self.columnas_list.clear()
        for c in table.columns:
            ctype = c.type.replace(" PRIMARY KEY", " [PK]")
            item = QListWidgetItem(f"# {c.name}  [{ctype}]")
            item.setData(Qt.ItemDataRole.UserRole, c.name)
            item.setToolTip("Clic para inyectar en el editor")
            self.columnas_list.addItem(item)
        try:
            lines = [f"Tabla: {table.name}"] + [f"  • {c.name} — {c.type}" for c in table.columns]
            self.columnas_label.setText("\n".join(lines))
        except Exception:
            pass
        self._refresh_dump(table)

    def _on_columna_clicked(self, item: QListWidgetItem) -> None:
        col = item.data(Qt.ItemDataRole.UserRole)
        if col:
            tc = self.editor.textCursor()
            tc.insertText(col)
            self.editor.setTextCursor(tc)
            self.editor.setFocus()

    def _refresh_dump(self, table: Table) -> None:
        total = len(table.rows)
        ver = min(total, self.VISOR_MAX_FILAS)
        cols = [c.name for c in table.columns]
        types = [c.type.split()[0] for c in table.columns]
        headers = [f"{c} ::{t}" for c, t in zip(cols, types)]
        self.visor_tabla.clear()
        self.visor_tabla.setColumnCount(len(cols))
        self.visor_tabla.setHorizontalHeaderLabels(headers)
        self.visor_tabla.setRowCount(ver)
        for r, row in enumerate(table.rows[:ver]):
            for c, value in enumerate(row):
                self.visor_tabla.setItem(r, c, _item_grilla(value))
        ajustar_y_repartir(self.visor_tabla, headers, table.rows, self.MUESTRA_MEDICION)

    def _on_historial_clicked(self, item: QListWidgetItem) -> None:
        query = item.data(Qt.ItemDataRole.UserRole)
        if query:
            self.editor.setPlainText(query)
            self.ejecutar_consulta()

    # ------------------------------------------------------- ejecución

    def _toast(self, text: str) -> None:
        self.toast_msg = text.upper()
        self.exec_time.setText(self.toast_msg)
        self.mensaje_label.setText(text)

    def ejecutar_consulta(self) -> None:
        query = self.editor.toPlainText().strip()
        if not query:
            self._mostrar_error("EDITOR VACÍO: la consulta está vacía. Escribe una instrucción SQL para evaluar.", None)
            return
        t0 = time.perf_counter()
        result = self.engine.execute(query)
        ms = (time.perf_counter() - t0) * 1000
        if result.error:
            self.exec_time.setText("FALLO_EJEC // ERROR")
            self._mostrar_error(result.error, None)
            self._refresh_status()
            return
        self.exec_time.setText(f"T_EJEC: {ms:.2f} ms // ESTADO: 200 OK")
        self._mostrar_resultado(result.columns, result.rows)
        self.exec_time.setText(f"T_EJEC: {ms:.2f} ms // ESTADO: 200 OK")
        self._refresh_status()
        if query:
            self._add_to_historial(query)

    def _mostrar_resultado(self, columns: list[str], rows: list[list]) -> None:
        self.error_box.setVisible(False)
        if not columns:
            self._limpiar_resultado()
            return
        self.empty_state.setVisible(False)
        self.resultado_tabla.setVisible(True)
        self.row_badge.setVisible(True)
        total = len(rows)
        ver = min(total, self.RESULTADO_MAX_FILAS)
        self.row_badge.setText(f"{total} {'FILA' if total == 1 else 'FILAS'}")
        self.resultado_tabla.clear()
        self.resultado_tabla.setColumnCount(len(columns))
        self.resultado_tabla.setHorizontalHeaderLabels(columns)
        self.resultado_tabla.setRowCount(ver)
        for r, row in enumerate(rows[:ver]):
            for c, value in enumerate(row):
                self.resultado_tabla.setItem(r, c, _item_grilla(value))
        ajustar_y_repartir(self.resultado_tabla, columns, rows, self.MUESTRA_MEDICION)
        self._ultimo_resultado = (columns, rows)
        self._toast(
            f"CONSULTA OK: {total} FILA(S)"
            + (f" (MOSTRANDO {ver})" if total > ver else "")
        )

    def _mostrar_error(self, text: str, _hint: str | None) -> None:
        self.empty_state.setVisible(False)
        self.resultado_tabla.setVisible(False)
        self.row_badge.setVisible(False)
        self.error_box.setVisible(True)
        self.error_text.setText(text)
        self.error_hint.setText("Verifica el nombre en ESQUEMA DE TABLAS: tablas y columnas exactas.")
        self.mensaje_label.setText(text)

    def _add_to_historial(self, query: str) -> None:
        self.historial.insert(0, query)
        self.historial = self.historial[:20]
        item = QListWidgetItem(f"> {query.replace(chr(10), ' ')}")
        item.setData(Qt.ItemDataRole.UserRole, query)
        item.setToolTip(datetime.now().strftime("%H:%M:%S") + " — clic para [RECARGAR]")
        self.historial_list.insertItem(0, item)
        while self.historial_list.count() > 20:
            self.historial_list.takeItem(self.historial_list.count() - 1)

    # ------------------------------------------------------- acciones

    def formatear_consulta(self) -> None:
        q = self.editor.toPlainText()
        if not q.strip():
            return
        self.editor.setPlainText(_formatear_sql(q).strip())
        self._toast("SINTAXIS SQL FORMATEADA AL ESTÁNDAR")

    def limpiar_editor(self) -> None:
        self.editor.clear()
        self.editor.setFocus()

    def limpiar_historial(self) -> None:
        self.historial = []
        self.historial_list.clear()

    def comparar_desde_historial(self) -> None:
        """Diálogo para elegir 2 consultas del historial y compararlas (CP-03)."""
        from core.comparar import comparar_consultas
        consultas = [
            self.historial_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.historial_list.count())
        ]
        consultas = [q for q in consultas if q]
        if len(consultas) < 2:
            self._toast("SE NECESITAN 2 CONSULTAS EN EL HISTORIAL")
            return
        from PySide6.QtWidgets import QListWidget
        dlg = QDialog(self)
        dlg.setWindowTitle("COMPARAR CONSULTAS")
        dlg.setMinimumWidth(520)
        dlg.setModal(True)
        lay = QVBoxLayout(dlg)
        lay.setSpacing(8)
        lab_a = QLabel("CONSULTA A:")
        lab_a.setObjectName("MutedLabel")
        lay.addWidget(lab_a)
        list_a = QListWidget()
        for q in consultas:
            list_a.addItem(str(q).replace("\n", " "))
        list_a.setCurrentRow(0)
        lay.addWidget(list_a)
        lab_b = QLabel("CONSULTA B:")
        lab_b.setObjectName("MutedLabel")
        lay.addWidget(lab_b)
        list_b = QListWidget()
        for q in consultas:
            list_b.addItem(str(q).replace("\n", " "))
        list_b.setCurrentRow(1 if len(consultas) > 1 else 0)
        lay.addWidget(list_b)
        row = QHBoxLayout()
        row.addStretch()
        eleccion: list[bool] = [False]
        btn_ok = QPushButton("COMPARAR")
        btn_ok.setObjectName("PrimaryBtn")
        btn_ok.clicked.connect(lambda: (eleccion.__setitem__(0, True), dlg.accept()))
        btn_no = QPushButton("CANCELAR")
        btn_no.setObjectName("GhostBtn")
        btn_no.clicked.connect(dlg.reject)
        row.addWidget(btn_ok)
        row.addWidget(btn_no)
        lay.addLayout(row)
        if dlg.exec() != QDialog.DialogCode.Accepted or not eleccion[0]:
            return
        ia = list_a.currentRow()
        ib = list_b.currentRow()
        if ia < 0 or ib < 0:
            return
        veredicto = comparar_consultas(consultas[ia], consultas[ib], self.engine)
        lineas = [veredicto["resumen"]] + [f"{donde}: {fila}" for donde, fila in veredicto["diff"]]
        titulo = "IGUALES" if veredicto["iguales"] else "DIFERENTES"
        self._toast(f"COMPARACIÓN: {titulo} // {veredicto['resumen']}")
        _show_custom_dialog(self, f"COMPARACIÓN: {titulo}", "\n".join(lineas))

    def ver_select_all(self) -> None:
        item = self.tabla_list.currentItem()
        if not item:
            return
        name = item.data(Qt.ItemDataRole.UserRole)
        self.editor.setPlainText(f'SELECT * FROM "{name}";')
        self.ejecutar_consulta()

    def restablecer_datos(self) -> None:
        if not self.cargar_preset("ejemplo_tienda.json", silencioso=True):
            self._refresh_tabla_list()
            self._refresh_briefing()
            self._set_default_query()
            self._limpiar_resultado()
        self._toast("BASE DE DATOS RESTABLECIDA AL ESTADO INICIAL")

    def copiar_consulta(self) -> None:
        query = self.editor.toPlainText().strip()
        if not query:
            self._toast("EDITOR VACÍO // NINGUNA CONSULTA EN EL EDITOR")
            return
        from PySide6.QtGui import QGuiApplication
        titulo = self.ejercicio.titulo or "ejercicio"
        QGuiApplication.clipboard().setText(
            f'Hola IA, resolví el ejercicio "{titulo}".\n'
            f"Esta es mi consulta SQL:\n\n```sql\n{query}\n```\n\n"
            "¿Es la solución óptima o cómo puedo mejorarla?"
        )
        self._toast("CONSULTA EXPORTADA PARA IA // COPIADA")

    def copiar_especificacion(self) -> None:
        """Copia título + enunciado + columnas objetivo (+ orden) — CE-01/CE-02."""
        if not (self.ejercicio.titulo or self.ejercicio.enunciado):
            self._toast("SIN EJERCICIO // NADA QUE COPIAR")
            return
        from PySide6.QtGui import QGuiApplication
        lineas = [self.ejercicio.titulo, "", self.ejercicio.enunciado]
        cols = self._expected_columns()
        if cols:
            lineas += ["", "COLUMNAS OBJETIVO: " + ", ".join(cols)]
        if self.sort_label.text().strip():
            lineas.append(self.sort_label.text().strip())
        QGuiApplication.clipboard().setText("\n".join(lineas))
        self._toast("ESPECIFICACIÓN COPIADA")

    def mostrar_formato_json(self) -> None:
        dlg = QDialog(self)
        dlg.setWindowTitle("ESPECIFICACIÓN DE PROTOCOLO JSON PARA IA")
        dlg.resize(560, 480)
        lay = QVBoxLayout(dlg)
        lay.setSpacing(8)
        title = QLabel("ESPECIFICACIÓN DE PROTOCOLO JSON PARA IA")
        title.setObjectName("PanelTitle")
        lay.addWidget(title)
        info = QLabel("Pega este prompt en IA para generar retos compatibles. Guarda el .json e impórtalo con [CARGAR EJERCICIO].")
        info.setObjectName("MutedLabel")
        info.setWordWrap(True)
        lay.addWidget(info)
        editor = QPlainTextEdit()
        editor.setReadOnly(True)
        editor.setPlainText(CLAUDE_PROMPT)
        self.claude_prompt_text = editor.toPlainText()  # compat
        lay.addWidget(editor, stretch=1)
        row = QHBoxLayout()
        row.addStretch()
        btn_copy = QPushButton("COPIAR PLANTILLA")
        btn_copy.setObjectName("PrimaryBtn")
        btn_close = QPushButton("CERRAR")
        btn_close.setObjectName("GhostBtn")
        row.addWidget(btn_copy)
        row.addWidget(btn_close)
        lay.addLayout(row)
        toast_lbl = QLabel("")
        toast_lbl.setObjectName("ToastLabel")
        lay.addWidget(toast_lbl)

        def _copy() -> None:
            from PySide6.QtGui import QGuiApplication
            QGuiApplication.clipboard().setText(CLAUDE_PROMPT)
            toast_lbl.setText("✓ PLANTILLA COPIADA AL PORTAPAPELES")
            btn_copy.setText("✓ COPIADA")
            btn_copy.setEnabled(False)

            def _clear_toast() -> None:
                try:
                    toast_lbl.setText("")
                except RuntimeError:
                    pass  # diálogo ya cerrado: nada que limpiar

            QTimer.singleShot(2400, _clear_toast)
            self._toast("PLANTILLA COPIADA AL PORTAPAPELES")

        btn_copy.clicked.connect(_copy)
        btn_close.clicked.connect(dlg.accept)
        dlg.exec()

    # ------------------------------------------------------- sesiones

    def exportar_db(self) -> None:
        """Vuelca la memoria a un .db real (ED-02/ED-03)."""
        if not self.engine.tables:
            self._toast("SIN DATOS // NADA QUE EXPORTAR")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "EXPORTAR BASE A DB", "base.db", "SQLite (*.db)"
        )
        if not path:
            return
        try:
            self.engine.exportar_db(path)
        except (OSError, ValueError) as exc:
            _show_custom_dialog(self, "ERROR AL EXPORTAR", f"No se pudo escribir:\n{exc}")
            return
        n = len(self.engine.tables)
        self._toast(f"BASE EXPORTADA: {os.path.basename(path)} ({n} TABLAS)")

    def guardar_sesion(self) -> None:
        if not self.engine.tables:
            _show_custom_dialog(self, "Sin datos", "No hay tablas cargadas para guardar.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Guardar sesión", "sesion.json", "Archivo JSON (*.json)")
        if not path:
            return
        payload = {
            "ejercicio": {
                "titulo": self.ejercicio.titulo,
                "enunciado": self.ejercicio.enunciado,
                "pista": self.ejercicio.pista,
                "dificultad": self.ejercicio.dificultad,
                "default_query": self.ejercicio.default_query,
            },
            "tablas": [
                {
                    "nombre": t.name,
                    "columnas": [{"nombre": c.name, "tipo": c.type} for c in t.columns],
                    "filas": t.rows,
                }
                for t in self.engine.tables.values()
            ],
            "historial": self.historial[:20],
        }
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        self._toast(f"SESIÓN GUARDADA: {os.path.basename(path)} EN DISCO")

    def cargar_sesion(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Cargar sesión", "", "Archivo JSON (*.json)")
        if not path:
            return
        if not self._confirmar_archivo_grande([path]):
            return
        result = load_file(path)
        self._aplicar_resultado(result, path)
        try:
            with open(path, "r", encoding="utf-8-sig") as f:
                data = json.load(f)
            hist = data.get("historial", [])
            self.historial = []
            self.historial_list.clear()
            for query in reversed(hist):
                self._add_to_historial(query)
        except Exception:
            pass

    def closeEvent(self, event: QCloseEvent) -> None:
        self.engine.close()
        event.accept()

    def exportar_resultado_csv(self) -> None:
        """Guarda el último resultado completo en .csv (EX-01..EX-06)."""
        columns, rows = self._ultimo_resultado
        if not columns or not rows:
            self._toast("SIN RESULTADOS // NADA QUE EXPORTAR")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "EXPORTAR RESULTADO A CSV", "resultado.csv", "CSV (*.csv)"
        )
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(columns)
                for row in rows:
                    w.writerow(["NULL" if v is None else v for v in row])
        except OSError as exc:
            _show_custom_dialog(self, "ERROR AL EXPORTAR", f"No se pudo escribir:\n{exc}")
            return
        self._toast(f"RESULTADO EXPORTADO: {os.path.basename(path)} ({len(rows)} FILAS)")

    def exportar_resultado_excel(self) -> None:
        """Guarda el último resultado completo en .xlsx real (XL-01..XL-04)."""
        columns, rows = self._ultimo_resultado
        if not columns or not rows:
            self._toast("SIN RESULTADOS // NADA QUE EXPORTAR")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "EXPORTAR RESULTADO A EXCEL", "resultado.xlsx", "Excel (*.xlsx)"
        )
        if not path:
            return
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "resultado"
            ws.append(list(columns))
            for row in rows:
                ws.append(["NULL" if v is None else v for v in row])
            for c in range(1, len(columns) + 1):
                ancho = 10
                for r in range(1, min(len(rows) + 2, 102)):
                    val = ws.cell(r, c).value
                    if val is not None:
                        ancho = max(ancho, min(len(str(val)) + 2, 50))
                ws.column_dimensions[ws.cell(1, c).column_letter].width = ancho
            wb.save(path)
        except OSError as exc:
            _show_custom_dialog(self, "ERROR AL EXPORTAR", f"No se pudo escribir:\n{exc}")
            return
        self._toast(f"RESULTADO EXPORTADO: {os.path.basename(path)} ({len(rows)} FILAS)")

    def graficar_resultado(self) -> None:
        """Abre el diálogo de gráfico del último resultado (GR-01/GR-02)."""
        from ui.graficos import DialogoGrafico, datos_para_grafico, puede_graficar
        columns, rows = self._ultimo_resultado
        puntos = datos_para_grafico(columns, rows) if puede_graficar(columns, rows) else None
        if puntos is None:
            self._toast("SE NECESITAN 2 COLUMNAS (ETIQUETA, VALOR) // NADA QUE GRAFICAR")
            return
        etiquetas, valores = puntos
        titulo = "GRÁFICO"
        if len(rows) > 200:
            from ui.graficos import MAX_PUNTOS
            titulo = f"GRÁFICO (PRIMERAS {MAX_PUNTOS} FILAS)"
        dlg = DialogoGrafico(self, titulo, etiquetas, valores)
        dlg.exec()

    def _limpiar_resultado(self) -> None:
        self._ultimo_resultado = ([], [])
        self.resultado_tabla.clear()
        self.resultado_tabla.setRowCount(0)
        self.resultado_tabla.setColumnCount(0)
        self.resultado_tabla.setVisible(False)
        self.row_badge.setVisible(False)
        self.error_box.setVisible(False)
        self.empty_state.setVisible(True)
        self.exec_time.setText("EN ESPERA")
