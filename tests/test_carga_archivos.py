"""Spec carga-tablas-archivos — multi-selección de archivos + carpeta intacta."""
from __future__ import annotations

from pathlib import Path

from PySide6.QtWidgets import QDialog, QFileDialog, QPushButton


def _make_xlsx(path: Path, rows: list[list]):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for r in rows:
        ws.append(r)
    wb.save(str(path))


def test_combinar_csv_y_xlsx(tmp_path):
    """CA-02: fusiona tablas de varios archivos."""
    from core.session_loader import combinar_resultados, load_file

    a = tmp_path / "a.csv"
    a.write_text("id,n\n1,2\n", encoding="utf-8")
    b = tmp_path / "b.xlsx"
    _make_xlsx(b, [["id", "x"], [1, "q"]])
    merged, reemplazadas = combinar_resultados([load_file(str(a)), load_file(str(b))])
    assert merged.ok, merged.errors
    assert sorted(t.name for t in merged.tables) == ["a", "b"]
    assert reemplazadas == []


def test_combinar_last_wins(tmp_path):
    """CA-03: mismo nombre de tabla → gana el último + aviso."""
    from core.session_loader import combinar_resultados, load_file

    d1 = tmp_path / "uno"
    d1.mkdir()
    d2 = tmp_path / "dos"
    d2.mkdir()
    (d1 / "a.csv").write_text("id\n1\n", encoding="utf-8")
    _make_xlsx(d2 / "a.xlsx", [["id"], [1], [2]])
    merged, reemplazadas = combinar_resultados(
        [load_file(str(d1 / "a.csv")), load_file(str(d2 / "a.xlsx"))]
    )
    assert merged.ok
    assert len(merged.tables) == 1
    assert merged.tables[0].name == "a"
    assert len(merged.tables[0].rows) == 2
    assert reemplazadas == ["a"]


def test_combinar_invalido_no_bloquea_validos(tmp_path):
    """CA-04: archivo inválido entre válidos → válidos cargan + error con nombre."""
    from core.session_loader import combinar_resultados, load_file

    ok = tmp_path / "ok.csv"
    ok.write_text("id\n1\n", encoding="utf-8")
    malo = tmp_path / "malo.csv"
    malo.write_text("id\n", encoding="utf-8")  # solo cabecera
    merged, _ = combinar_resultados([load_file(str(ok)), load_file(str(malo))])
    assert merged.ok
    assert [t.name for t in merged.tables] == ["ok"]
    assert any("malo.csv" in e for e in merged.errors)


def test_combinar_todos_invalidos(tmp_path):
    """Todos inválidos → ok=False."""
    from core.session_loader import combinar_resultados, load_file

    m = tmp_path / "m.csv"
    m.write_text("id\n", encoding="utf-8")
    merged, _ = combinar_resultados([load_file(str(m))])
    assert not merged.ok
    assert merged.tables == []


def test_combinar_conserva_ejercicio(tmp_path):
    """El primer ejercicio no vacío se conserva."""
    from core.session_loader import combinar_resultados, load_file

    j = tmp_path / "e.json"
    j.write_text(
        '{"ejercicio": {"titulo": "T", "enunciado": "E", "pista": "", "dificultad": "Fácil"},'
        ' "tablas": [{"nombre": "t", "columnas": [{"nombre": "id", "tipo": "INTEGER"}], "filas": [[1]]}]}',
        encoding="utf-8",
    )
    c = tmp_path / "c.csv"
    c.write_text("id\n1\n", encoding="utf-8")
    merged, _ = combinar_resultados([load_file(str(j)), load_file(str(c))])
    assert merged.ok
    assert merged.ejercicio is not None and merged.ejercicio.titulo == "T"
    assert sorted(t.name for t in merged.tables) == ["c", "t"]


def test_ui_cargar_archivos_multi(app, monkeypatch, tmp_path):
    """CA-01/02: el slot carga lo que devuelve getOpenFileNames."""
    a = tmp_path / "a.csv"
    a.write_text("id\n1\n2\n", encoding="utf-8")
    b = tmp_path / "b.csv"
    b.write_text("id\n9\n", encoding="utf-8")
    monkeypatch.setattr(
        QFileDialog, "getOpenFileNames",
        lambda *args, **kwargs: ([str(a), str(b)], "Tablas (*.csv *.xlsx *.xls)"),
    )
    app.cargar_tablas_archivos()
    assert set(app.engine.table_names()) >= {"a", "b"}


def test_ui_cargar_archivos_cancelar_no_crashea(app, monkeypatch):
    """CA-06: cancelar → retorno silencioso."""
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *a, **k: ([], ""))
    before = set(app.engine.table_names())
    app.cargar_tablas_archivos()
    assert set(app.engine.table_names()) == before


def test_ui_elegir_modo_archivos(app, monkeypatch):
    """CA-01: mini-diálogo [ARCHIVOS] devuelve 'archivos'."""

    def fake_exec(self):
        for btn in self.findChildren(QPushButton):
            if btn.text() == "ARCHIVOS":
                btn.click()
                break
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(QDialog, "exec", fake_exec)
    assert app._elegir_modo_carga() == "archivos"


def test_ui_elegir_modo_carpeta(app, monkeypatch):
    """CA-05: mini-diálogo [CARPETA] devuelve 'carpeta'."""

    def fake_exec(self):
        for btn in self.findChildren(QPushButton):
            if btn.text() == "CARPETA":
                btn.click()
                break
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(QDialog, "exec", fake_exec)
    assert app._elegir_modo_carga() == "carpeta"


def test_ui_cargar_tablas_despacha_archivos(app, monkeypatch):
    """El botón CARGAR TABLAS despacha según el modo elegido."""
    llamadas = []
    monkeypatch.setattr(app, "_elegir_modo_carga", lambda: "archivos")
    monkeypatch.setattr(app, "cargar_tablas_archivos", lambda: llamadas.append("archivos"))
    monkeypatch.setattr(app, "cargar_csv_carpeta", lambda: llamadas.append("carpeta"))
    app.cargar_tablas()
    assert llamadas == ["archivos"]


def test_ui_cargar_tablas_despacha_carpeta(app, monkeypatch):
    monkeypatch.setattr(app, "_elegir_modo_carga", lambda: "carpeta")
    llamadas = []
    monkeypatch.setattr(app, "cargar_tablas_archivos", lambda: llamadas.append("archivos"))
    monkeypatch.setattr(app, "cargar_csv_carpeta", lambda: llamadas.append("carpeta"))
    app.cargar_tablas()
    assert llamadas == ["carpeta"]
