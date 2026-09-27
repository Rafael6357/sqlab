"""Spec carga-tipos-pg — BOOLEAN/DATE/TIMESTAMP en carga (CT)."""
from __future__ import annotations

import os

import pytest

from core.pg_engine import PG_BIN_DIR
from core.session_loader import _infer_type, load_file

REQUIERE_PG = not (PG_BIN_DIR and os.path.isfile(os.path.join(PG_BIN_DIR, "initdb.exe")))
salta_sin_pg = pytest.mark.skipif(REQUIERE_PG, reason="sin binarios PG vendoreados")


def test_ct01_inferencia():
    """CT-01: bool/fecha/timestamp; 't' y fechas imposibles son TEXT."""
    assert _infer_type("true") == "BOOLEAN"
    assert _infer_type("FALSE") == "BOOLEAN"
    assert _infer_type("yes") == "BOOLEAN"
    assert _infer_type("t") == "TEXT"
    assert _infer_type("2024-05-01") == "DATE"
    assert _infer_type("2024-05-01 10:00:00") == "TIMESTAMP"
    assert _infer_type("2024-13-45") == "TEXT"
    assert _infer_type("NULL") == "TEXT"
    assert _infer_type("12") == "INTEGER"


def test_ct02_csv_bool_fecha(tmp_path):
    """CT-02: CSV infiere BOOLEAN/DATE y convierte a bool real."""
    p = tmp_path / "t.csv"
    p.write_text("vip,alta\ntrue,2024-05-01\nFALSE,2024-06-02\n", encoding="utf-8")
    res = load_file(str(p))
    assert res.ok, res.errors
    t = res.tables[0]
    assert [c.type for c in t.columns] == ["BOOLEAN", "DATE"]
    assert t.rows == [[True, "2024-05-01"], [False, "2024-06-02"]]


def test_ct02_json_bool(tmp_path):
    """CT-02: JSON con BOOLEAN declarado convierte strings; NA → None."""
    import json
    p = tmp_path / "e.json"
    payload = {"tablas": [{"nombre": "t",
                           "columnas": [{"nombre": "v", "tipo": "BOOLEAN"}],
                           "filas": [["true"], ["NA"], ["false"]]}]}
    p.write_text(json.dumps(payload), encoding="utf-8")
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].rows == [[True], [None], [False]]


def test_ct02_excel_bool(tmp_path):
    """CT-02: Excel con bool nativo + texto 'true' en columna bool."""
    import openpyxl
    p = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["vip"])
    ws.append([True])
    ws.append(["true"])
    wb.save(str(p))
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].columns[0].type == "BOOLEAN"
    assert res.tables[0].rows == [[True], [True]]


@salta_sin_pg
def test_ct03_e2e_pg(tmp_path):
    """CT-03: CSV → PGEngine → filtros bool y fecha (DATE vuelve como date)."""
    import datetime
    from core.pg_engine import PGEngine
    p = tmp_path / "t.csv"
    p.write_text("vip,alta\ntrue,2024-05-01\nfalse,2024-06-02\n", encoding="utf-8")
    res = load_file(str(p))
    assert res.ok, res.errors
    eng = PGEngine(base_dir=str(tmp_path))
    try:
        assert eng.load_tables(res.tables) == []
        r = eng.execute("SELECT alta FROM t WHERE vip = TRUE")
        assert r.ok and r.rows == [[datetime.date(2024, 5, 1)]]
        r = eng.execute("SELECT COUNT(*) FROM t WHERE alta > DATE '2024-05-15'")
        assert r.ok and r.rows == [[1]]
    finally:
        eng.close()
