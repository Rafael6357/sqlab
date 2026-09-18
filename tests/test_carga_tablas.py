"""Spec carga-tablas-excel-csv — carga robusta CSV + XLSX + XLS."""
from __future__ import annotations

from pathlib import Path



def _make_xlsx(path: Path, rows: list[list]):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    for r in rows:
        ws.append(r)
    wb.save(str(path))


def _make_xls(path: Path, rows: list[list]):
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Hoja1")
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.write(i, j, v)
    wb.save(str(path))


def _make_xlsx_multi(path: Path, sheet_rows: dict[str, list[list]]):
    import openpyxl
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheet_rows.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(title=name)
        for r in rows:
            ws.append(r)
    wb.save(str(path))


# EC-02: CSV con extensión mayúsculas
def test_csv_uppercase_ext(tmp_path):
    from core.session_loader import load_file, load_tablas_folder

    folder = tmp_path / "upper"
    folder.mkdir()
    (folder / "DATOS.CSV").write_text("id,nombre\n1,Ana\n2,Luis\n", encoding="utf-8")
    res = load_tablas_folder(str(folder))
    assert res.ok, res.errors
    assert len(res.tables) == 1
    assert res.tables[0].name == "DATOS"
    assert len(res.tables[0].rows) == 2
    # load_file también debe aceptar mayúsculas directo
    res2 = load_file(str(folder / "DATOS.CSV"))
    assert res2.ok, res2.errors


def test_csv_bom(tmp_path):
    """EC-02: CSV con BOM utf-8-sig no contamina la primera cabecera."""
    from core.session_loader import load_file

    p = tmp_path / "bom.csv"
    # Escribir con BOM explícito
    p.write_bytes(b"\xef\xbb\xbfid,nombre\n1,Ana\n2,Beto\n")
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].columns[0].name == "id"
    assert res.tables[0].columns[1].name == "nombre"


def test_csv_semicolon(tmp_path):
    """EC-02: CSV delimitado por ; con auto-detección."""
    from core.session_loader import load_file

    p = tmp_path / "semi.csv"
    p.write_text("id;nombre;edad\n1;Ana;30\n2;Luis;25\n", encoding="utf-8")
    res = load_file(str(p))
    assert res.ok, res.errors
    assert [c.name for c in res.tables[0].columns] == ["id", "nombre", "edad"]
    assert len(res.tables[0].rows) == 2


def test_csv_only_header_error(tmp_path):
    """EC-06: solo cabecera → error explícito sin crash."""
    from core.session_loader import load_file

    p = tmp_path / "solo.csv"
    p.write_text("id,nombre\n", encoding="utf-8")
    res = load_file(str(p))
    assert not res.ok
    assert any("vacío o solo tiene cabecera" in e for e in res.errors)


# EC-03
def test_xlsx_first_sheet(tmp_path):
    from core.session_loader import load_file

    p = tmp_path / "ventas.xlsx"
    _make_xlsx(p, [["id", "producto", "precio"], [1, "Laptop", 1200], [2, "Mouse", 25]])
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].name == "ventas"
    assert [c.name for c in res.tables[0].columns] == ["id", "producto", "precio"]
    assert len(res.tables[0].rows) == 2


def test_xlsx_uppercase_ext(tmp_path):
    from core.session_loader import load_file

    p = tmp_path / "VENTAS.XLSX"
    _make_xlsx(p, [["id", "nombre"], [1, "Ana"]])
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].name == "VENTAS"


def test_xlsx_multi_sheet_solo_primera(tmp_path):
    """Límites: solo 1ª hoja."""
    from core.session_loader import load_file

    p = tmp_path / "multi.xlsx"
    _make_xlsx_multi(p, {"Hoja1": [["id", "a"], [1, "x"]], "Hoja2": [["id", "b"], [9, "y"]]})
    res = load_file(str(p))
    assert res.ok, res.errors
    assert [c.name for c in res.tables[0].columns] == ["id", "a"]
    assert res.tables[0].rows == [[1, "x"]]


# EC-04
def test_xls_legacy(tmp_path):
    from core.session_loader import load_file

    p = tmp_path / "listado.xls"
    _make_xls(p, [["id", "nombre"], [1, "Ana"], [2, "Luis"]])
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].name == "listado"
    assert [c.name for c in res.tables[0].columns] == ["id", "nombre"]
    assert len(res.tables[0].rows) == 2


# EC-05
def test_mixed_folder_csv_xlsx_xls(tmp_path):
    from core.session_loader import load_tablas_folder

    folder = tmp_path / "mix"
    folder.mkdir()
    (folder / "a.csv").write_text("id,nombre\n1,Ana\n", encoding="utf-8")
    _make_xlsx(folder / "b.XLSX", [["id", "prod"], [1, "X"]])
    _make_xls(folder / "c.Xls", [["id", "cosa"], [1, "Y"]])
    res = load_tablas_folder(str(folder))
    assert res.ok, res.errors
    assert len(res.tables) == 3
    names = {t.name for t in res.tables}
    assert names == {"a", "b", "c"}


# EC-06
def test_empty_folder_error_mentions_all_exts(tmp_path):
    from core.session_loader import load_tablas_folder

    folder = tmp_path / "vacia"
    folder.mkdir()
    res = load_tablas_folder(str(folder))
    assert not res.ok
    msg = "\n".join(res.errors)
    assert "*.csv" in msg and "*.xlsx" in msg and "*.xls" in msg


def test_csv_cabecera_vacia_normalizada(tmp_path):
    """Edge: cabecera vacía → col{n}."""
    from core.session_loader import load_file

    p = tmp_path / "vacia.csv"
    p.write_text("id,,edad\n1,Ana,30\n", encoding="utf-8")
    res = load_file(str(p))
    assert res.ok, res.errors
    assert res.tables[0].columns[1].name.startswith("col")


def test_xlsx_celda_vacia_none(tmp_path):
    from core.session_loader import load_file

    p = tmp_path / "huecos.xlsx"
    _make_xlsx(p, [["id", "nombre", "edad"], [1, "Ana", None], [2, None, 25]])
    res = load_file(str(p))
    assert res.ok, res.errors
    # celdas vacías → None
    assert res.tables[0].rows[0][2] is None
    assert res.tables[0].rows[1][1] is None


def test_alias_load_csv_folder_sigue(tmp_path):
    """Compat: el alias histórico load_csv_folder sigue existiendo."""
    import core.session_loader as sl

    assert hasattr(sl, "load_csv_folder")
    folder = tmp_path / "alias"
    folder.mkdir()
    (folder / "a.csv").write_text("id,n\n1,2\n", encoding="utf-8")
    res = sl.load_csv_folder(str(folder))
    assert res.ok


def test_ui_boton_tablas_y_tooltip(app):
    """EC-01: el botón ahora dice TABLAS y su tooltip documenta el formato."""
    assert hasattr(app, "btn_csv")
    assert "TABLAS" in app.btn_csv.text().upper()
    tip = app.btn_csv.toolTip().lower()
    assert "*.csv" in tip
    assert "*.xlsx" in tip or "xlsx" in tip
    assert "*.xls" in tip or "xls" in tip


def test_load_file_formato_no_soportado(tmp_path):
    from core.session_loader import load_file

    p = tmp_path / "a.txt"
    p.write_text("hola", encoding="utf-8")
    res = load_file(str(p))
    assert not res.ok
    assert "Formato no soportado" in "\n".join(res.errors)
